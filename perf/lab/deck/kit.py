"""Utilidades de estilo para el deck (python-pptx)."""
import copy
from lxml import etree
from pptx import Presentation
from pptx.chart.data import XyChartData, CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

W, H = 13.333, 7.5
FONT = 'Segoe UI'
MONO = 'Consolas'


def rgb(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


C = {
    'bg': '#fcfcfb', 'plane': '#f3f2ee', 'ink': '#0b0b0b', 'ink2': '#52514e', 'muted': '#898781',
    'grid': '#e1e0d9', 'axis': '#c3c2b7', 'border': '#dcdbd4',
    's1': '#2a78d6', 's2': '#eb6834', 's3': '#1baf7a', 's7': '#4a3aa7',
    'critical': '#d03b3b', 'good': '#0ca30c', 'warning': '#fab219', 'goodtext': '#006300',
    'blue100': '#cde2fb', 'blue250': '#86b6ef', 'blue600': '#184f95', 'orange100': '#fbe1d6',
    'code': '#f3f2ee',
}


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.layout = self.prs.slide_layouts[6]
        self.n = 0

    def slide(self, title=None, kicker=None, notes=None):
        s = self.prs.slides.add_slide(self.layout)
        self.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(C['bg'])
        if kicker:
            text(s, 0.6, 0.38, 10, 0.3, kicker.upper(), size=11, color=C['s1'], bold=True, spacing=1.5)
        if title:
            text(s, 0.6, 0.62, 12.1, 0.8, title, size=28, color=C['ink'], bold=True)
        if self.n > 1:
            text(s, 12.1, 7.05, 0.8, 0.3, str(self.n), size=10, color=C['muted'], align=PP_ALIGN.RIGHT)
            text(s, 0.6, 7.05, 8, 0.3, 'Taller de pruebas de carga', size=10, color=C['muted'])
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self, path):
        self.prs.save(path)


def text(s, x, y, w, h, content, size=16, color=None, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, line=1.1, italic=False):
    """content: str o lista de párrafos; cada párrafo str o lista de runs (texto, {opts})."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        runs = para if isinstance(para, list) else [(para, {})]
        for r_text, opts in runs:
            r = p.add_run()
            r.text = r_text
            f = r.font
            f.name = opts.get('font', font)
            f.size = Pt(opts.get('size', size))
            f.bold = opts.get('bold', bold)
            f.italic = opts.get('italic', italic)
            f.color.rgb = rgb(opts.get('color', color or C['ink']))
            if spacing:
                r._r.get_or_add_rPr().set('spc', str(int(spacing * 100)))
        if isinstance(para, list) and para and para[0][1].get('space_after') is not None:
            p.space_after = Pt(para[0][1]['space_after'])
    return tb


def para_space(tb, pts):
    for p in tb.text_frame.paragraphs:
        p.space_after = Pt(pts)


def box(s, x, y, w, h, fill=None, line=None, radius=0.06, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.margin_left = sh.text_frame.margin_right = Inches(0.12)
    return sh


def shape_text(sh, content, size=14, color=None, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=FONT):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [(para, {})]
        for r_text, opts in runs:
            r = p.add_run()
            r.text = r_text
            r.font.name = opts.get('font', font)
            r.font.size = Pt(opts.get('size', size))
            r.font.bold = opts.get('bold', bold)
            r.font.color.rgb = rgb(opts.get('color', color or C['ink']))
    return sh


def line(s, x1, y1, x2, y2, color=None, width=1.25, dash=False, arrow_end=False, arrow_start=False):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color or C['ink2'])
    ln.line.width = Pt(width)
    lnxml = ln.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(lnxml, qn('a:prstDash'))
        d.set('val', 'dash')
    if arrow_end:
        t = etree.SubElement(lnxml, qn('a:tailEnd'))
        t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
    if arrow_start:
        t = etree.SubElement(lnxml, qn('a:headEnd'))
        t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
    return ln


def code(s, x, y, w, h, src, size=12, highlight=None):
    """highlight: dict linea_idx -> color de texto."""
    b = box(s, x, y, w, h, fill=C['code'], radius=0.03)
    tb = s.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(w - 0.4), Inches(h - 0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln_ in enumerate(src.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05
        r = p.add_run()
        r.text = ln_ if ln_ else ' '
        r.font.name = MONO
        r.font.size = Pt(size)
        col = (highlight or {}).get(i)
        r.font.color.rgb = rgb(col or C['ink'])
        r.font.bold = bool(col)
    return b


# ------------------------------------------------------------------ charts

def _plot_layout(chart, x, y, w, h):
    """Fija el área interna del gráfico (fracciones del marco) para poder anotar encima."""
    plotArea = chart._chartSpace.find('.//' + qn('c:plotArea'))
    old = plotArea.find(qn('c:layout'))
    if old is not None:
        plotArea.remove(old)
    layout = etree.SubElement(plotArea, qn('c:layout'))
    plotArea.remove(layout)
    plotArea.insert(0, layout)
    ml = etree.SubElement(layout, qn('c:manualLayout'))
    for tag, val in [('c:layoutTarget', 'inner'), ('c:xMode', 'edge'), ('c:yMode', 'edge'),
                     ('c:x', x), ('c:y', y), ('c:w', w), ('c:h', h)]:
        e = etree.SubElement(ml, qn(tag))
        e.set('val', str(val))


class Plot:
    """Gráfico nativo con área de trazado fija y mapeo datos->pulgadas para anotaciones."""

    def __init__(self, s, x, y, w, h, kind, data, xr=None, yr=None, xlog=False, ylog=False,
                 inner=(0.09, 0.05, 0.86, 0.78), legend=False, xtitle=None, ytitle=None,
                 yfmt='General', xfmt='General', ymajor=None, xmajor=None, font_size=11, gap=60):
        self.s, self.x, self.y, self.w, self.h = s, x, y, w, h
        self.inner = inner
        self.xr, self.yr, self.xlog, self.ylog = xr, yr, xlog, ylog
        gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
        self.chart = ch = gf.chart
        ch.has_title = False
        ch.font.name = FONT
        ch.font.size = Pt(font_size)
        ch.font.color.rgb = rgb(C['ink2'])
        ch.has_legend = legend
        if legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(font_size)
        _plot_layout(ch, *inner)
        self.kind = kind
        va = ch.value_axis
        self._axis(va, yr, ylog, ymajor, yfmt, grid=True)
        if kind in (XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, XL_CHART_TYPE.XY_SCATTER_LINES,
                    XL_CHART_TYPE.XY_SCATTER):
            ca = ch.category_axis
            self._axis(ca, xr, xlog, xmajor, xfmt, grid=False)
        else:
            ca = ch.category_axis
            ca.format.line.color.rgb = rgb(C['axis'])
            ca.major_tick_mark = XL_TICK_MARK.NONE
            ca.tick_labels.font.size = Pt(font_size)
            ca.tick_labels.font.color.rgb = rgb(C['ink2'])
            plot = ch.plots[0]
            try:
                plot.gap_width = gap
                plot.overlap = -10
            except Exception:
                pass
        if xtitle:
            self._title(ca, xtitle, font_size)
        if ytitle:
            self._title(va, ytitle, font_size)

    def _axis(self, ax, r, log, major, fmt, grid):
        ax.format.line.color.rgb = rgb(C['axis'])
        ax.major_tick_mark = XL_TICK_MARK.NONE
        ax.minor_tick_mark = XL_TICK_MARK.NONE
        ax.tick_labels.font.color.rgb = rgb(C['ink2'])
        ax.tick_labels.number_format = fmt
        ax.tick_labels.number_format_is_linked = False
        if r:
            ax.minimum_scale, ax.maximum_scale = r
        if major:
            ax.major_unit = major
        ax.has_major_gridlines = grid
        if grid:
            ax.major_gridlines.format.line.color.rgb = rgb(C['grid'])
            ax.major_gridlines.format.line.width = Pt(0.75)
        else:
            ax.has_major_gridlines = False
        if log:
            from pptx.enum.chart import XL_AXIS_CROSSES
            ax.crosses = XL_AXIS_CROSSES.MINIMUM
            scaling = ax._element.find(qn('c:scaling'))
            lb = etree.Element(qn('c:logBase'))
            lb.set('val', '10')
            scaling.insert(0, lb)

    def _title(self, ax, t, size):
        ax.has_title = True
        tf = ax.axis_title.text_frame
        tf.text = t
        f = tf.paragraphs[0].runs[0].font
        f.size = Pt(size)
        f.bold = False
        f.color.rgb = rgb(C['ink2'])
        f.name = FONT

    def style_series(self, idx, color, width=2.25, marker=False, dash=False, fill=None, smooth=False, no_line=False, marker_size=7):
        ser = self.chart.plots[0].series[idx]
        if self.kind in (XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.COLUMN_STACKED):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(fill or color)
            ser.format.line.color.rgb = rgb(C['bg'])
            ser.format.line.width = Pt(1.5)
            return ser
        if no_line:
            ser.format.line.fill.background()
        else:
            ser.format.line.color.rgb = rgb(color)
            ser.format.line.width = Pt(width)
        ser.smooth = smooth
        if dash:
            lnxml = ser.format.line._get_or_add_ln()
            d = etree.SubElement(lnxml, qn('a:prstDash'))
            d.set('val', 'dash')
        from pptx.enum.chart import XL_MARKER_STYLE
        if marker:
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = marker_size
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = rgb(color)
            ser.marker.format.line.color.rgb = rgb(C['bg'])
            ser.marker.format.line.width = Pt(1.5)
        else:
            ser.marker.style = XL_MARKER_STYLE.NONE
        return ser

    # mapeo datos -> pulgadas en la diapositiva
    def px(self, v):
        import math
        lo, hi = self.xr
        if self.xlog:
            f = (math.log10(v) - math.log10(lo)) / (math.log10(hi) - math.log10(lo))
        else:
            f = (v - lo) / (hi - lo)
        return self.x + self.w * (self.inner[0] + self.inner[2] * f)

    def py(self, v):
        import math
        lo, hi = self.yr
        if self.ylog:
            f = (math.log10(v) - math.log10(lo)) / (math.log10(hi) - math.log10(lo))
        else:
            f = (v - lo) / (hi - lo)
        return self.y + self.h * (self.inner[1] + self.inner[3] * (1 - f))

    def cat_x(self, i, n):
        """centro de la categoría i de n en un gráfico de categorías."""
        return self.x + self.w * (self.inner[0] + self.inner[2] * (i + 0.5) / n)

    @property
    def left(self):
        return self.x + self.w * self.inner[0]

    @property
    def right(self):
        return self.x + self.w * (self.inner[0] + self.inner[2])

    @property
    def top(self):
        return self.y + self.h * self.inner[1]

    @property
    def bottom(self):
        return self.y + self.h * (self.inner[1] + self.inner[3])


def xy(series):
    """series: lista de (nombre, [(x,y),...])"""
    d = XyChartData()
    for name, pts in series:
        s = d.add_series(name)
        for a, b in pts:
            s.add_data_point(a, b)
    return d


def cat(categories, series, fmt=None):
    d = CategoryChartData(number_format=fmt) if fmt else CategoryChartData()
    d.categories = categories
    for name, vals in series:
        d.add_series(name, vals)
    return d


def pill(s, x, y, label, color, fill, size=11, w=None):
    w = w or (0.12 * len(label) * size / 11 + 0.35)
    b = box(s, x, y, w, 0.32, fill=fill, radius=0.5)
    shape_text(b, label, size=size, color=color, bold=True)
    return b


def dot(s, cx, cy, d, color):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    o.fill.solid()
    o.fill.fore_color.rgb = rgb(color)
    o.line.color.rgb = rgb(C['bg'])
    o.line.width = Pt(1.5)
    o.shadow.inherit = False
    return o


def hide_legend_entry(chart, idx):
    leg = chart._chartSpace.find('.//' + qn('c:legend'))
    e = etree.SubElement(leg, qn('c:legendEntry'))
    i = etree.SubElement(e, qn('c:idx')); i.set('val', str(idx))
    dl = etree.SubElement(e, qn('c:delete')); dl.set('val', '1')
    leg.remove(e)
    leg.insert(1, e)
